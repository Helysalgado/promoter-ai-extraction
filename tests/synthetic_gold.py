"""Synthetic gold-workbook helpers for evaluator-only tests.

Workbooks are written under pytest ``tmp_path`` only.  They mimic the current
workset layout (sheet ``Hoja1``, headers on row 4) without copying private
gold values.
"""
from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook

from promoter_ai_extraction.models import (
    EvidenceItem,
    ExtractedValue,
    ExtractionRun,
    Property,
    PropertyResult,
    ScientificStatus,
)
from promoter_ai_extraction.persistence import (
    PersistedPredictionRef,
    PredictionStore,
    VerifiedPersistedPrediction,
)

GOLD_HEADERS: tuple[str, ...] = (
    "Fila_origen",
    "ID_promotor",
    "Nombre_promotor",
    "Sinonimo_gen_en_este_paper",
    "ID_paper",
    "Propiedad",
    "Valor_RegulonDB",
    "Sin_dato_en_RegulonDB",
    "Modalidad_origen",
    "Valor_verificado_manualmente",
    "GT_para_referencia",
    "Año_confirmado",
    "Técnica_confirmada_manualmente",
    "Evidencia",
)


def write_synthetic_gold(path: Path, rows: list[dict[str, Any]]) -> Path:
    """Write a synthetic Hoja1 workbook with headers on row 4."""
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Hoja1"
    sheet["A1"] = "synthetic evaluator fixture"
    for column, header in enumerate(GOLD_HEADERS, start=1):
        sheet.cell(4, column, header)
    for row_index, row in enumerate(rows):
        for key, value in row.items():
            sheet.cell(5 + row_index, GOLD_HEADERS.index(key) + 1, value)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path


def make_run(run_id: str = "eval-run-001", paper_id: str = "PMC12345") -> ExtractionRun:
    evidence = EvidenceItem(
        fragment="The TSS was at -42.",
        segment_id="seg:0001",
        source_type="body_text",
        location="line 1",
    )
    extracted = PropertyResult(
        paper_id=paper_id,
        promoter_name="lacZp1",
        property=Property.TSS,
        status=ScientificStatus.EXTRACTED,
        values=(
            ExtractedValue("-42", "-42", None, None, (evidence,)),
        ),
        candidate_values=(),
        evidence=(evidence,),
        abstention_reason=None,
    )

    def abstained(prop: Property) -> PropertyResult:
        return PropertyResult(
            paper_id=paper_id,
            promoter_name="lacZp1",
            property=prop,
            status=ScientificStatus.NOT_FOUND,
            values=(),
            candidate_values=(),
            evidence=(),
            abstention_reason="Not found in the synthetic document.",
        )

    return ExtractionRun(
        run_id=run_id,
        paper_id=paper_id,
        promoter_name="lacZp1",
        tss=extracted,
        caja_10=abstained(Property.CAJA_10),
        caja_35=abstained(Property.CAJA_35),
        sigma=abstained(Property.FACTOR_SIGMA),
    )


def persist_verified(
    tmp_path: Path, run: ExtractionRun | None = None
) -> VerifiedPersistedPrediction:
    store = PredictionStore(tmp_path / "predictions")
    saved = store.save(run or make_run())
    assert isinstance(saved, PersistedPredictionRef)
    verified = store.load_verified(saved)
    assert isinstance(verified, VerifiedPersistedPrediction)
    return verified
